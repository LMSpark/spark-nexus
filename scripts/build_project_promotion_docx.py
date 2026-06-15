from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "交付物"
OUT_DOCX = OUT_DIR / "SPARK_Nexus城市物业治理中枢_项目推介书_从理念到落地.docx"

FONT_CN = "Microsoft YaHei"
FONT_EN = "Calibri"

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
TEAL = RGBColor(0, 105, 92)
GOLD = RGBColor(122, 90, 0)
RED = RGBColor(153, 27, 27)
GRAY = RGBColor(82, 94, 108)
LIGHT_FILL = "F4F6F9"
TABLE_FILL = "E8EEF5"
MUTED_FILL = "F2F4F7"
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)


def rgb_hex(rgb):
    return f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"


def set_run_font(run, size=None, color=None, bold=None, italic=None, font=FONT_CN):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), FONT_EN)
    rfonts.set(qn("w:hAnsi"), FONT_EN)
    rfonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_paragraph(p, before=0, after=8, line=1.333, align=None, keep_next=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        p.alignment = align
    if keep_next:
        fmt.keep_with_next = True


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    tbl.insert(0, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for name, size, color, before, after in [
        ("Title", 24, NAVY, 0, 10),
        ("Subtitle", 13, GRAY, 0, 10),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.208

    for list_style in ["List Bullet", "List Number"]:
        style = styles[list_style]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("SPARK Nexus 城市物业治理中枢项目推介书")
    set_run_font(run, size=9, color=GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("领码科技 | Page ")
    set_run_font(run, size=9, color=GRAY)
    add_page_number(fp)


def add_para(doc, text="", size=11, color=BLACK, bold=False, italic=False, align=None, before=0, after=8, line=1.333):
    p = doc.add_paragraph()
    style_paragraph(p, before=before, after=after, line=line, align=align)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_kicker(doc, text):
    p = add_para(doc, text, size=10, color=GOLD, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, line=1.15)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run)
    if level == 1:
        p.paragraph_format.page_break_before = False
    p.paragraph_format.keep_with_next = True
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_paragraph(p, after=4, line=1.208)
        run = p.add_run(item)
        set_run_font(run, size=10.5, color=BLACK)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_paragraph(p, after=4, line=1.208)
        run = p.add_run(item)
        set_run_font(run, size=10.5, color=BLACK)


def add_callout(doc, title, body, fill=LIGHT_FILL, accent=TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360], indent_dxa=120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    style_paragraph(p, before=2, after=2, line=1.25)
    r1 = p.add_run(title + "：")
    set_run_font(r1, size=10.5, color=accent, bold=True)
    r2 = p.add_run(body)
    set_run_font(r2, size=10.5, color=BLACK)
    add_para(doc, "", after=4)


def add_table(doc, headers, rows, widths_dxa, header_fill=TABLE_FILL, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths_dxa, indent_dxa=120)
    hdr = table.rows[0]
    repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0, line=1.15)
        r = p.add_run(h)
        set_run_font(r, size=font_size, color=NAVY, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cell = cells[i]
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(val)) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, after=0, line=1.2)
            r = p.add_run(str(val))
            set_run_font(r, size=font_size, color=BLACK)
    set_table_geometry(table, widths_dxa, indent_dxa=120)
    add_para(doc, "", after=4)
    return table


def add_image(doc, rel_path, caption, width=6.1):
    path = ROOT / rel_path
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    c = add_para(doc, caption, size=9, color=GRAY, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, line=1.15)
    return c


def add_cover(doc):
    add_para(doc, "领码科技", size=12, color=GRAY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8, line=1.15)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("SPARK Nexus 城市物业治理中枢")
    set_run_font(r, size=23, color=NAVY, bold=True)
    p2 = doc.add_paragraph(style="Subtitle")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(4)
    r = p2.add_run("项目推介书 | 政策落地与试点推广版")
    set_run_font(r, size=16, color=GRAY, bold=True)
    add_para(
        doc,
        "把人民至上落到小区，把群众路线落到服务，把公共收益晒在阳光下。",
        size=12,
        color=TEAL,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=6,
        line=1.15,
    )
    add_para(
        doc,
        "SPARK Nexus 不是再造一个物业收费系统，而是以人民至上为价值底座、以小区为最小治理单元，把政府监管、居委会准入、业委会决策、物业服务、银行资金和本地服务商生态接入同一套数字秩序。",
        size=9.8,
        color=GRAY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=10,
        line=1.15,
    )

    hero = ROOT / "web-admin" / "src" / "assets" / "community-governance-hero.png"
    if hero.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(str(hero), width=Inches(5.55))
        p.paragraph_format.space_after = Pt(8)

    rows = [
        ["适用对象", "政府/街道、居委会、业委会、物业公司、银行、本地服务商"],
        ["文档定位", "面向领导汇报、试点申报、商务沟通和实施启动的统一材料"],
        ["核心主题", "把人民至上落到小区，把公共收益管清楚，把民生服务做扎实"],
        ["版本日期", "2026年6月13日"],
    ]
    add_table(doc, ["项目", "说明"], rows, [1700, 7660], header_fill=MUTED_FILL, font_size=8.6)

    doc.add_page_break()


def add_toc(doc):
    add_heading(doc, "目录", 1)
    entries = [
        "1. 一句话价值：把小区治理从“糊涂账”变成“明白账”",
        "2. 文化理念与建设理念：以小区为中心的五方共治",
        "3. 政府政策框架与全国类似探索",
        "4. 痛点诊断：为什么小区需要一套“明白账”",
        "5. 产品定位：城市物业治理操作系统",
        "6. 平台总体架构",
        "7. 核心功能体系",
        "8. 关键业务闭环",
        "9. 技术架构与信创适配",
        "10. 数据安全、审计与合规体系",
        "11. 实施路径与交付计划",
        "12. 运营落地机制",
        "13. 商业模式与投入产出",
        "14. 风险控制与应对",
        "15. 验收标准与交付物",
        "16. 推广路径：先跑通一个样板，再复制一座城市",
        "附录 A. 功能清单",
        "附录 B. 参考资料与材料来源",
    ]
    add_bullets(doc, entries)
    doc.add_page_break()


def build_doc():
    OUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "SPARK Nexus城市物业治理中枢项目推介书"
    doc.core_properties.author = "领码科技"
    doc.core_properties.subject = "项目推介、实施方案、商业模式、落地路径"
    doc.core_properties.keywords = "物业管理,智慧社区,公共收益,银行对账,基层治理,SPARK Nexus"

    add_cover(doc)
    add_toc(doc)

    add_heading(doc, "1. 一句话价值：把小区治理从“糊涂账”变成“明白账”", 1)
    add_para(
        doc,
        "一句话概括：SPARK Nexus 把小区里“公共收益怎么管、民生服务谁负责、支出审批谁监督、居民诉求谁闭环”全部搬到同一张数字台账上，让党的群众路线有入口、基层治理有抓手、民生福祉有证据、群众监督有回音。项目以“小区主档”为中心，以居委会准入和数据授权为秩序，以物业收费、公共收益、银行对账、业主服务、监管审计为业务闭环，形成一套可试点、可验收、可复制、可持续运营的平台方案。",
    )
    add_callout(
        doc,
        "传播主张",
        "人民至上、阳光透明、合规闭环、群众有感。上了是样板，不上是隐患；早一天上线，群众早一天受益，基层早一天主动；晚一天透明，多一天糊涂账、多一分信任流失。",
        fill="EDF7F6",
        accent=TEAL,
    )
    add_table(
        doc,
        ["对象", "想要什么", "上了得到什么", "不上的代价"],
        [
            ["政府/街道", "人民至上、群众有感、治理有力", "一套样板、一组数据、一包证据，把党的初心使命落到小区日常", "公共收益不透明，群众信任就会流失，基层矛盾就会累积。"],
            ["物业", "收费更顺、投诉更少、续约更稳", "账单好收、解释有据、工单闭环，物业从挨骂变成有底气", "继续靠嘴解释，业主不信任会吞掉所有服务努力。"],
            ["银行", "入口、资金、流水、场景", "从缴费通道升级为小区资金入口，拿到监管账户和社区金融场景", "只做通道就会被替代，先占小区资金入口才有未来。"],
            ["业主", "看见钱、管住事、说话有用", "公共收益可查、支出可追、服务可评，业主从旁观者变共同管理人", "钱不清，信任就回不来；事不闭环，矛盾只会越积越大。"],
        ],
        [1450, 2700, 3050, 2160],
        font_size=8.25,
    )
    add_callout(
        doc,
        "强记忆句",
        "小区虽小，连着民心；账本虽薄，照见作风；服务虽细，关系党心。公共收益不晒出来，信任就会继续流失；服务过程不留痕，责任就会继续扯皮；监管证据不在线，风险就会继续潜伏。",
        fill="FFF8E8",
        accent=GOLD,
    )

    add_heading(doc, "2. 文化理念与建设理念：以小区为中心的五方共治", 1)
    add_para(
        doc,
        "一座城市的治理能力，最终会落到一个个小区里。业主是否信任物业，公共收益是否公开，报修投诉是否闭环，老人儿童是否被看见，银行资金是否安全，政府监管是否有据，都是党的群众工作在基层治理中的真实考题。SPARK Nexus 的文化理念，就是把这些看似琐碎的日常事务，变成一套有规则、有温度、有证据的社区运行机制。",
    )
    add_callout(
        doc,
        "文化主张",
        "人民至上不是一句口号，而是一套能被群众感受到的制度安排；共治不是一次会议，而是一条链路；服务不是一个工单，而是居民可感知的改善。平台的价值不止于上线系统，而是把基层治理中难以说清、难以追责、难以协同的事项，沉淀为可看见、可核验、可持续运营的城市社区数字秩序。",
        fill="EDF7F6",
        accent=TEAL,
    )
    add_callout(
        doc,
        "信念底座",
        "中国共产党的根本宗旨是全心全意为人民服务。落到小区，就是群众的钱有人管、群众的事有人办、群众的诉求有回音、群众的监督有入口。SPARK Nexus 的政治表达不是大屏炫技，而是把人民至上、群众路线、依法治理和数字治理变成基层干部每天能用、群众每天能感到的治理工具。",
        fill="FFF1F0",
        accent=RED,
    )
    add_table(
        doc,
        ["理念维度", "表达", "项目落点"],
        [
            ["使命", "让小区治理从“说不清”走向“查得到”，从“各管一段”走向“一链到底”。", "小区主档、关系台账、授权边界、验收中心和审计留痕"],
            ["愿景", "让每个小区都有一本阳光账、一套协同链、一个群众服务入口。", "公共收益、物业收费、银行对账、报修投诉和监管预警闭环"],
            ["价值观", "阳光透明、责任共担、居民有感、长期主义。", "账务公开、实名表决、服务评价、信用评分和持续运营指标"],
            ["设计原则", "少造概念，多建闭环；少做展示，多留证据；少靠人情，多靠规则。", "准入审批、数据授权、回调验签、防重放、导出审计和交付验收"],
        ],
        [1550, 4150, 3660],
        font_size=8.8,
    )
    add_image(doc, "宣传片输出/preview-sheet.png", "图 1 项目理念宣传片抽帧：从公共收益治理升级为多租户、多主体、可运营的城市物业治理平台。", width=5.8)

    add_heading(doc, "2.1 建设理念：小区不是末梢，而是治理现场", 2)
    add_para(
        doc,
        "项目的核心理念不是把多个既有系统简单拼在一起，而是把小区定义为治理和运营的最小闭环单元。一个小区对应一套基础主档、一套参与主体关系、一套数据授权边界、一套资金闭环和一套服务评价体系。所有业务围绕小区展开，所有主体通过关系和授权进入小区，而不是各自维护一套孤立数据。",
    )
    add_heading(doc, "2.2 一套秩序：准入、授权、留痕", 2)
    add_bullets(
        doc,
        [
            "准入：小区备案、物业服务、银行服务、本地服务商进入小区，均通过平台审批形成证据。",
            "授权：政府、物业、银行、业委会、服务商只能访问其被授权的小区和数据范围。",
            "留痕：资金、审批、导出、回调、敏感字段核验、服务评价等关键动作形成审计链。",
            "协同：居委会、业委会、物业、银行和监管部门在同一个小区视图中分工协作。",
        ],
    )
    add_heading(doc, "2.3 五方价值主张", 2)
    add_table(
        doc,
        ["主体", "核心诉求", "平台价值"],
        [
            ["政府/街道", "群众有感、治理有力、风险可控、验收有据", "监管驾驶舱、风险预警、民生服务台账、验收中心、审计导出"],
            ["居委会", "有治理抓手、有审批入口、有居民触达", "小区准入、党建台账、网格事项、关爱走访、微信治理通知"],
            ["业委会/业主", "公共收益透明、重大支出可表决、服务可评价", "账务公开、实名投票、支出审批、投诉报修、消息回执"],
            ["物业公司", "收费率、服务效率、财务规范和业主满意度", "账单收费、报修投诉、电子票据、财务凭证、公示公告"],
            ["银行", "缴费场景、监管账户、对账和社区金融入口", "银行适配器、流水对账、放款回执、回调验签、数据交换"],
            ["本地服务商", "可信进入小区、获取服务订单、接受评价监管", "资质入驻、居委会审批、服务评价、投诉清退、结算扩展"],
        ],
        [1500, 2700, 5160],
        font_size=9,
    )

    add_heading(doc, "3. 政府政策框架与全国类似探索", 1)
    add_para(
        doc,
        "政策的方向已经很清楚：基层治理要下沉到社区，社区服务要连接到居民，公共收益要回到阳光下，智慧社区要从展示屏走向日常治理。SPARK Nexus 要讲的不是“我们有什么功能”，而是“国家政策为什么需要这样的数字底座、地方试点为什么需要这样的闭环工具、城市复制为什么需要这样的标准平台”。",
    )
    add_callout(
        doc,
        "政策判断",
        "政策正在把三个关键词推到台前：基层治理、阳光物业、公共收益。谁能把小区账户管清楚、把居民服务办顺畅、把监管证据留完整，谁就能成为下一轮智慧社区和物业治理升级的关键入口。SPARK Nexus 应顺势定位为政策落地工具，而不是单纯的物业收费系统。",
        fill="EDF7F6",
        accent=TEAL,
    )
    add_heading(doc, "3.1 国家政策框架：基层治理、智慧社区与完整社区", 2)
    add_table(
        doc,
        ["政策方向", "核心要求", "项目承接方式"],
        [
            ["基层治理现代化", "强化党建引领，推动乡镇、街道、社区组织对基层事务形成统筹协调和协商治理能力。", "以街道/居委会为治理入口，沉淀小区关系、事项流转、监督台账和验收证据。"],
            ["城乡社区服务体系", "提升社区服务、风险防范、矛盾调解、便民惠民和公共服务能力。", "把公告、报修、投诉、问卷、投票、关爱走访和居民触达放到统一入口。"],
            ["智慧社区建设", "用数字技术整合社区服务资源，形成线上线下融合的治理和服务新形态。", "建设管理端、业主端、监管端和数据交换能力，打通物业、银行、政府和居民数据。"],
            ["完整社区建设", "以居民委员会辖区为基本单元，补齐服务设施和生活服务短板，形成可复制样板。", "以试点小区和街道为单元推进，先跑通治理、资金、服务、监管和运营闭环。"],
        ],
        [1850, 4050, 3460],
        font_size=8.5,
    )
    add_heading(doc, "3.2 依法治理依据：业主共有权与公共收益透明化", 2)
    add_para(
        doc,
        "《民法典》确立业主对共有部分的共有和共同管理权，并规定建设单位、物业服务企业或者其他管理人利用业主共有部分产生的收入，在扣除合理成本之后属于业主共有。公共收益管理的关键由此转化为四件事：收益能不能准确归集，账目能不能单独核算，使用能不能共同决定，监管和审计能不能留痕追溯。",
    )
    add_bullets(
        doc,
        [
            "公共收益不是物业公司的附属收入，而是业主共有权益的数字化管理对象。",
            "物业服务、业委会决策、居委会监督、街道监管和银行账户应形成互相校验的治理链条。",
            "湖北、福建、河南等地已陆续出台住宅小区公共收益管理办法或试行办法，细化范围、账户、公示、审计、监督和责任。",
            "湖北政策解读还指出，公共收益制度不完善、管理不规范、账目不公开、监管不到位已经成为民生关切，亟需通过规则和平台共同治理。",
        ],
    )
    doc.add_page_break()
    add_heading(doc, "3.3 全国类似探索：从平台到制度的可复制样本", 2)
    add_table(
        doc,
        ["地区/来源", "已公开做法", "对本项目的启示"],
        [
            ["惠州", "智慧物业服务平台集成公共收益监管、物业服务信息公示、企业档案、业主诉求等功能。", "项目应把公共收益、服务监督和诉求处置放在同一平台，不割裂成多个入口。"],
            ["海南", "省住建厅公开征求规范住宅小区公共收益管理指导意见，关注收益范围、专项账户、公示、审计、监督和会计核算附件。", "海南样本说明公共收益正在从原则要求走向可操作表单，系统应内置账户、台账、公示表和审计证据。"],
            ["连云港", "试点住宅小区公共收益账户共管制度，提升资金透明度和使用效率。", "银行账户共管、流水对账、使用审批和公示留痕是公共收益落地的关键抓手。"],
            ["七台河", "智慧社区信息化平台整合多部门政务数据，服务基层治理数字化和社区工作减负。", "政府侧价值不仅是看大屏，更是减少重复采集、强化数据底座和一站式事项处理。"],
            ["上海美好家园案例", "通过智慧小区建设、党建引领和业主监督提升物业治理效率，公共收益反哺小区治理。", "项目应把技术、业主参与和公共收益使用效果连接起来，形成居民可感知的治理成果。"],
            ["新华社观察", "多地探索专户管理、透明核算、线上公开、业主决策和阶段性审计，让公共收益“晒在阳光下”。", "系统应内置专户、凭证、公示、表决、审计和监督闭环，减少线下扯皮空间。"],
        ],
        [1800, 4050, 3510],
        font_size=7.8,
    )
    add_heading(doc, "3.4 市场机会：谁先掌握小区入口，谁先形成城市样板", 2)
    add_para(
        doc,
        "传统物业 SaaS 有收费和工单，却缺少政府背书与公共收益监管；社区治理平台有网格和事件，却缺少支付、对账、票据和持续收入；银行缴费平台有支付结算，却缺少小区主档、物业授权和居民服务闭环。SPARK Nexus 的机会，是把三类系统都想做却没有做透的“小区入口”真正打通，形成治理准入、物业运营、银行资金和社区生态的复合型平台。",
    )

    add_heading(doc, "4. 痛点诊断：为什么小区需要一套“明白账”", 1)
    add_table(
        doc,
        ["痛点", "典型表现", "平台解法"],
        [
            ["钱不明", "公共收益、物业费、支出审批、银行流水分散管理，业主只看到结果，看不到过程。", "专户、账单、流水、凭证、审批和公示统一留痕，监管端可穿透查看。"],
            ["事不闭环", "报修、投诉、公告、投票、问卷各走各的渠道，过程散、结果散、责任散。", "围绕小区主档建立事项链路，提交、处理、评价、导出全流程留证。"],
            ["人不信任", "业主不信物业，物业怕解释，业委会难组织，居委会缺抓手。", "账务公开、实名投票、消息回执、服务评价和多方协同台账重建信任。"],
            ["银行只是通道", "银行只做支付入口，缺少监管账户、对账和场景金融延展。", "银行服务配置、多银行适配、监管账户、对账差异闭环、放款回执和数据交换。"],
            ["验收说不清", "项目落地后难以证明功能真实使用、业务真实闭环、风险真实下降。", "验收中心按数据自动核验功能状态，导出证据标签、缺口摘要和整改动作。"],
            ["生态进不来", "本地服务商靠熟人或群消息进入小区，缺准入、评价和退出机制。", "服务商资质入驻，居委会审批绑定小区，服务评价、投诉、暂停和清退留痕。"],
        ],
        [1700, 3400, 4260],
        font_size=8.8,
    )

    add_heading(doc, "5. 产品定位：城市物业治理操作系统", 1)
    add_para(
        doc,
        "SPARK Nexus 的定位可以更直接地表达为“城市物业治理操作系统”：向上承接政府和街道监管要求，向中间支撑居委会和业委会治理，向下连接物业、银行、居民和本地服务商的日常业务。它不是一次性展示系统，而是一套能每天产生数据、每月支撑运营、每年接受验收的城市社区数字底座。",
    )
    add_heading(doc, "5.1 总体目标", 2)
    add_bullets(
        doc,
        [
            "一张小区主档：把社区、楼栋、房屋、业主、机构和服务关系统一起来。",
            "一本阳光账本：把物业费、公共收益、银行流水、审批凭证和公示结果统一起来。",
            "一个居民入口：把缴费、报修、投诉、公告、投票、问卷和评价统一起来。",
            "一套监管证据：把风险预警、信用评分、导出审计、验收中心和数据交换统一起来。",
            "一条商业路径：从政府试点切入，用物业和银行形成长期收入，用服务商生态打开增长空间。",
        ],
    )
    add_heading(doc, "5.2 首期试点建议目标", 2)
    add_table(
        doc,
        ["类别", "建议指标", "验收口径"],
        [
            ["覆盖范围", "1个街道或区县、1-3个居委会、10-30个小区", "平台小区主档和关系台账可查"],
            ["物业运营", "2-5家物业公司、物业费/停车费/水电费账单上线", "线上缴费率、欠费台账、导出审计"],
            ["银行合作", "1-2家银行接入代收、对账或监管账户能力", "流水入库、对账差异闭环、回调验签"],
            ["居民服务", "业主端缴费、报修、投诉、投票、公告消息可用", "工单评价、消息回执、投票记录"],
            ["监管验收", "监管驾驶舱、风险预警、验收中心、数据交换可演示", "自动核验结果、证据导出文件"],
        ],
        [1700, 3800, 3860],
        font_size=9,
    )

    add_heading(doc, "6. 平台总体架构", 1)
    add_para(
        doc,
        "平台采用“三端一体 + 多主体协同 + 数据授权隔离”的总体架构。管理端服务政府、街道、居委会、业委会、物业、银行和平台运营方；业主端以 H5/小程序承接居民高频操作；后台服务统一承载权限、业务、审计、外部接口和数据交换。",
    )
    add_table(
        doc,
        ["层级", "组成", "说明"],
        [
            ["用户触达层", "管理端、业主端 H5/小程序、监管大屏", "面向不同角色提供差异化入口和菜单能力"],
            ["业务能力层", "注册准入、物业运营、资金监管、公共收益、报修投诉、投票问卷、党建网格", "围绕小区主档组织业务闭环"],
            ["资金与接口层", "微信/支付宝、银行直连、短信、实名认证、税务开票、对象存储", "外部能力按适配器配置，可 dev 模拟也可真实对接"],
            ["数据治理层", "租户、小区、房屋、业主、账单、流水、审批、凭证、审计、导出", "按租户、小区、数据范围和动作做授权隔离"],
            ["运维合规层", "验收中心、安全体检、部署就绪检查、日志脱敏、回调防重放", "支撑项目验收、上线检查和持续运维"],
        ],
        [1650, 3500, 4210],
        font_size=9,
    )
    add_image(doc, "outputs/admin-home-check.png", "图 2 管理端首页与登录入口示意，覆盖政府监管、物业运营、业委会治理和银行协同。", width=6.1)

    add_heading(doc, "7. 核心功能体系", 1)
    add_heading(doc, "7.1 政府监管与街道工作台", 2)
    add_bullets(
        doc,
        [
            "监管驾驶舱：辖区小区、公共收益、缴费率、收支趋势、风险预警、待审批支出。",
            "专题分析：资金异常、缴费风险、审批超时、信用评价等专题切片。",
            "信用评价：按缴费率、未关闭预警、欠费账单、审计链完整性等因子生成小区信用评分。",
            "数据交换：面向政府/银行/物业生成数据交换包，记录校验摘要、记录数和构建日志。",
            "验收中心：自动核验系统功能落地状态，展示缺口摘要、证据标签和下一步动作。",
        ],
    )
    add_heading(doc, "7.2 居委会与业委会治理", 2)
    add_bullets(
        doc,
        [
            "小区备案和服务商准入：物业、银行、本地服务商进入小区必须形成审批证据。",
            "党建和网格治理：党建活动、治理事项、关爱走访、资源联动、居民通知。",
            "公共收益决策：重大支出触发实名投票和多级审批，银行放款与审批结果联动。",
            "账务公开：收支明细、公共收益账户、凭证、现金流量表和账簿对业主可公示。",
            "议事协商：投票、问卷、公告、消息回执和结果留痕，支撑小区共治。",
        ],
    )
    add_heading(doc, "7.3 物业运营系统", 2)
    add_bullets(
        doc,
        [
            "收费账单：按小区、楼栋、房屋、费用类型和账期生成账单，支持批量导出。",
            "支付退款：微信/支付宝 dev 模式模拟，真实模式可替换渠道接口，退款生成负向流水和通知。",
            "报修投诉：住户提交、物业处理、完成评价，闭环留证。",
            "电子票据与税务适配：缴费后自动开具电子缴费票据，支持税务开票申请和红冲台账。",
            "财务凭证：收支单据自动制证、审核、记账、PDF 导出和账簿查询。",
        ],
    )
    add_heading(doc, "7.4 银行资金服务", 2)
    add_bullets(
        doc,
        [
            "小区银行服务配置：区分代收、退款、放款、监管账户、对账等服务能力。",
            "多银行适配：维护银行适配器 profile、签名算法、验签规则、对账模式和放款模式。",
            "流水与对账：银行流水导入、渠道账单同步、对账记录、差异明细和差异处理闭环。",
            "放款回执：公共收益审批通过后生成银行放款指令，回执回填流水号并保持幂等。",
            "回调安全：生产模式要求时间戳和 nonce，签名验真、防重放、重复业务幂等。",
        ],
    )
    add_heading(doc, "7.5 业主端小程序/H5", 2)
    add_bullets(
        doc,
        [
            "一站式服务：缴费、账务、公告、报修投诉、业主投票、问卷、统计台账。",
            "实名与多房屋：实名注册、房屋绑定申请、审核通过、默认房屋和房屋切换。",
            "动作隔离：问卷、投票、报修和评价按已审核房屋关系校验所属小区。",
            "消息回执：通知公告、缴费提醒、审批结果等支持未读/已读状态。",
        ],
    )
    add_image(doc, "outputs/owner-mobile-home-final.png", "图 3 业主端移动界面示意，覆盖缴费、账务、公告、投票、报修和房屋服务。", width=2.3)

    add_heading(doc, "8. 关键业务闭环", 1)
    add_heading(doc, "8.1 注册准入闭环", 2)
    add_numbered(
        doc,
        [
            "机构提交注册申请，平台或指定审核方核验资质。",
            "审核通过后自动创建租户、管理员账号和用户租户关系。",
            "物业、银行或服务商申请绑定小区，居委会或平台按规则审批。",
            "审批通过后生成小区关系和数据授权，菜单和接口按能力开放。",
            "停用关系时同步撤销授权，避免服务关系失效后仍保留数据权限。",
        ],
    )
    add_heading(doc, "8.2 物业收费闭环", 2)
    add_numbered(
        doc,
        [
            "物业配置收费标准，按账期生成账单并推送业主端。",
            "业主通过小程序/H5 缴费，支付渠道回调入账。",
            "系统自动更新账单状态，生成支付流水、渠道账单和电子缴费票据。",
            "银行流水同步后进入对账，差异项处理完成后对账单关闭。",
            "政府和物业按授权查看缴费率、欠费台账、流水和审计导出记录。",
        ],
    )
    add_heading(doc, "8.3 公共收益与支出审批闭环", 2)
    add_numbered(
        doc,
        [
            "公共收益进入业委会或小区公共收益专户，平台记录来源和凭证。",
            "支出申请关联发票、合同、验收单等材料，按工作流模板进入审批。",
            "达到阈值的重大支出触发业主实名投票，投票结果作为审批条件。",
            "审批通过后生成银行放款指令，银行回执回填流水；驳回或超时可触发退回。",
            "系统自动生成凭证、账簿、现金流量表和导出审计记录，支撑公示和监管。",
        ],
    )
    add_heading(doc, "8.4 报修投诉服务闭环", 2)
    add_numbered(
        doc,
        [
            "业主端提交报修、投诉或建议，系统校验房屋关系和小区权限。",
            "物业按工单类型派单处理，处理过程留痕并向业主推送状态。",
            "工单完成后必须形成服务评价证据，未完成评价可进入待办提醒。",
            "监管端可按小区、物业、时效、投诉率和满意度查看趋势。",
        ],
    )

    add_heading(doc, "9. 技术架构与信创适配", 1)
    add_para(
        doc,
        "当前工程已经具备 Spring Boot 后端、Vue 管理端、uni-app 业主端和本地 Docker 依赖。技术路线兼顾快速试点和后续信创扩展：业务先在 MySQL/Flyway/Redis/MinIO 等通用能力上跑通，再通过数据库方言、对象存储、国产操作系统和外部 SDK 适配进入生产环境。",
    )
    add_table(
        doc,
        ["模块", "当前实现", "生产/信创演进"],
        [
            ["后端", "Spring Boot 3 + Java 17 + REST API + Flyway", "微服务拆分、国产中间件、统一网关、统一身份"],
            ["管理端", "Vue 3 + Vite + TypeScript + Element Plus", "多角色菜单、监管大屏、国产浏览器兼容"],
            ["业主端", "uni-app，支持 H5 与微信小程序构建", "微信模板消息、真机联调、小程序发布流程"],
            ["数据层", "MySQL、Redis、MinIO、审计日志和数据导出", "达梦、人大金仓、openGauss、国产对象存储适配"],
            ["外部接口", "微信/支付宝/银行/短信/实名 dev 模式与配置台账", "真实 SDK/HTTP 签名客户端、回调验签、防重放"],
            ["安全运维", "部署就绪检查、安全体检、敏感字段审计", "等保测评、日志脱敏、密钥管理、持续安全扫描"],
        ],
        [1600, 3650, 4110],
        font_size=9,
    )
    add_heading(doc, "9.1 权限与数据模型", 2)
    add_para(
        doc,
        "系统从“用户角色 + 单小区”升级为“用户 + 当前租户 + 角色 + 小区 + 关系类型 + 数据范围 + 动作”的权限模型。机构作为租户，小区作为协作空间，关系表确定主体身份，授权表确定数据边界，注册审核流确定准入。",
    )
    add_bullets(
        doc,
        [
            "tenant：政府、物业、银行、平台、业委会等机构主体。",
            "user_tenant_relation：用户加入多个租户，并在不同租户拥有不同角色。",
            "tenant_community_relation：政府管辖、物业服务、银行代收、银行监管、业委会治理等小区关系。",
            "data_authorization：按小区、数据范围和动作授权 READ、WRITE、APPROVE、EXPORT、MANAGE。",
            "resident_house_relation：住户与多套房屋的实名绑定关系。",
        ],
    )

    add_heading(doc, "10. 数据安全、审计与合规体系", 1)
    add_para(
        doc,
        "平台涉及住户身份、房屋、账单、支付、银行流水、公共收益和审批材料，安全设计必须从权限边界、传输存储、接口回调、导出审计和运维配置五个层面闭环。项目推介阶段应强调能力边界：平台提供安全机制和合规留痕，但具体上线仍需结合部署单位安全制度、等保要求和正式接口规范执行。",
    )
    add_table(
        doc,
        ["安全域", "机制", "落地证据"],
        [
            ["身份与授权", "JWT、角色能力、租户切换、数据范围、详情防越权", "403 校验、能力菜单、授权台账导出"],
            ["敏感信息", "手机号、身份证号脱敏展示，核验行为写入审计", "sensitive_field_audit 与导出留档"],
            ["资金接口", "支付/银行回调验签、时间戳、nonce、防重放、幂等处理", "external_callback_receipt 与诊断工具"],
            ["数据导出", "所有账单、流水、报表、凭证、审计日志导出写入记录", "data_export_log 可追溯导出人、范围和行数"],
            ["部署就绪", "MySQL/Flyway、Redis、MinIO、CORS、密钥和生产配置风险检查", "DEPLOYMENT_READINESS 和安全体检清单"],
            ["审计链", "关键操作、审批、导出、敏感字段、回调、工作流事件留痕", "审计日志哈希链和模块化证据标签"],
        ],
        [1600, 4500, 3260],
        font_size=8.8,
    )

    add_heading(doc, "11. 实施路径与交付计划", 1)
    add_para(
        doc,
        "实施应采用“先试点、再复制、后生态”的路线。试点阶段不要贪大求全，而要把准入、收费、资金、服务、监管和验收六条闭环跑通，形成可演示、可运营、可审计的样板。",
    )
    add_table(
        doc,
        ["阶段", "周期", "关键任务", "交付成果"],
        [
            ["准备阶段", "第 0 周", "确定试点小区、参与主体、银行接口、数据范围和验收指标", "项目启动会、实施清单、账号角色表、数据模板"],
            ["基础上线", "第 1-2 周", "部署环境、导入小区/房屋/业主/机构，开通管理端和业主端", "可登录系统、基础档案、菜单权限、部署就绪报告"],
            ["业务闭环", "第 3-4 周", "账单缴费、报修投诉、公告消息、公共收益、银行流水和对账", "首轮业务数据、演示用例、导出审计记录"],
            ["监管验收", "第 5-6 周", "监管大屏、风险预警、信用评分、验收中心和数据交换包", "验收报告、缺口整改清单、演练证据包"],
            ["试运营", "第 7-12 周", "培训物业/居委会/银行，跟踪线上缴费率、工单闭环率和居民反馈", "试运营周报、KPI 看板、优化版本"],
            ["复制推广", "3-6 个月", "扩展小区、物业集团、银行分支和本地服务商生态", "区县/城市复制方案、商务合同模板、运营手册"],
        ],
        [1350, 1200, 4050, 2760],
        font_size=8.7,
    )
    add_heading(doc, "11.1 四周最小可落地计划", 2)
    add_numbered(
        doc,
        [
            "第 1 周：完成环境部署、基础档案导入、角色权限配置和系统可用性检查。",
            "第 2 周：完成注册审核、小区关系、数据授权和业主房屋绑定流程。",
            "第 3 周：完成账单、支付、银行流水、对账、报修投诉和公告消息闭环。",
            "第 4 周：完成公共收益、支出审批、业主投票、监管大屏、验收中心和导出证据。",
        ],
    )

    add_heading(doc, "12. 运营落地机制", 1)
    add_para(
        doc,
        "项目上线后成败不只取决于系统功能，还取决于谁来维护基础数据、谁来审核准入、谁来处理工单、谁来对账、谁来复盘指标。建议建立“街道/居委会牵头、物业运营、银行协同、平台支撑”的运营机制。",
    )
    add_table(
        doc,
        ["运营事项", "责任主体", "频率", "关键产出"],
        [
            ["小区与机构关系维护", "平台运营方/居委会", "按申请", "小区关系台账、数据授权记录"],
            ["账单与缴费运营", "物业公司", "月度", "账单明细、缴费率、欠费台账、电子票据"],
            ["银行对账", "物业/银行/平台", "日/周/月", "流水、渠道账单、对账差异处理结果"],
            ["公共收益公示", "业委会/物业/居委会", "月度/季度", "收支明细、审批记录、投票结果、公示记录"],
            ["监管复盘", "政府/街道/居委会", "月度", "风险预警、信用评分、治理事项闭环率"],
            ["系统运维", "平台运营方", "持续", "部署就绪检查、安全体检、数据备份和升级记录"],
            ["居民触达", "居委会/物业", "按事项", "公告通知、消息回执、问卷投票和服务评价"],
        ],
        [2400, 1900, 1200, 3860],
        font_size=8.8,
    )
    add_heading(doc, "12.1 培训与推广", 2)
    add_bullets(
        doc,
        [
            "政府/街道：重点培训监管驾驶舱、风险预警、验收中心、数据交换和审计导出。",
            "居委会：重点培训小区备案、服务商准入、党建网格、居民通知和公共收益监督。",
            "物业：重点培训收费配置、账单生成、支付退款、工单处理、票据和财务凭证。",
            "银行：重点培训银行服务配置、流水导入、对账差异、回调诊断和放款回执。",
            "业主：通过小程序二维码、物业公告、社区通知和线下培训完成首批绑定与缴费引导。",
        ],
    )

    add_heading(doc, "13. 商业模式与投入产出", 1)
    add_para(
        doc,
        "这类项目不能只靠一次性建设费。更稳的路径是：政府侧关注人民立场能否落地、群众获得感能否看见、基层风险能否前置化解；物业买的是收费率、服务效率和业主信任，银行买的是场景入口、资金沉淀和对账服务，服务商买的是合规进入小区的资格与订单机会。民生治理样板负责点火，物业与银行负责续航，服务商生态负责放大。",
    )
    add_table(
        doc,
        ["收入来源", "客户", "建议定价口径", "说明"],
        [
            ["物业 SaaS 年费", "物业公司/物业集团", "按项目、户数、模块收费；基础版可按 0.5-2 元/户/月或 2,000-8,000 元/月/项目测算", "直接服务收费、报修、公告、公示、票据和财务"],
            ["物业专业版", "中大型物业", "8-30 万元/年/公司", "公共收益、审批、财务凭证、电子票据、税务适配"],
            ["银行场景服务", "银行分支/总行部门", "10-50 万元/年/城市或区县", "多银行适配、对账、监管账户、放款、数据交换"],
            ["民生治理示范项目", "区县/街道", "20-100 万元/项目，含部署、培训、验收", "不作为长期主要收入，主要用于启动、验证和复制"],
            ["私有化/信创部署", "政府、国企物业、银行", "50-300 万元/项目", "受部署规模、接口复杂度、信创要求影响较大"],
            ["服务商生态", "本地生活服务商", "认证/上架费 + 订单服务费", "需坚持居委会审批和居民评价，避免无序商业化"],
        ],
        [1600, 1700, 3550, 2510],
        font_size=8.2,
    )
    add_heading(doc, "13.1 投入产出逻辑", 2)
    add_bullets(
        doc,
        [
            "对政府/街道：建设一套平台，形成群众有感的民生服务抓手、可推广的基层治理样板和可验收的治理证据。",
            "对物业：投入一个工具，换来收费率、服务效率、财务合规和业主信任的同步提升。",
            "对银行：投入一个场景，换来缴费交易、监管账户、对账服务、资金沉淀和社区金融入口。",
            "对平台：前期用样板换市场，后期用标准化复制降低边际交付成本。",
        ],
    )

    add_heading(doc, "14. 风险控制与应对", 1)
    add_table(
        doc,
        ["风险", "表现", "应对策略"],
        [
            ["多方利益复杂", "物业、业委会、居委会、银行诉求不一致", "用准入审批、关系表和授权表把权责边界写入系统"],
            ["示范项目周期长", "项目制推进慢，决策链条长", "政府侧围绕民生福祉、治理证据和风险防范做样板，长期收入靠物业和银行续费"],
            ["物业付费意愿不足", "小物业预算有限，重视短期效果", "用线上缴费率、欠费催缴、票据和服务评价证明经营价值"],
            ["银行接口复杂", "不同银行签名、对账、放款和回调差异大", "建立多银行适配器模板，先 dev 模式验证契约，再替换真实接口"],
            ["数据安全压力", "涉及住户身份、房屋、支付、银行流水", "租户隔离、脱敏、回调验签、防重放、导出审计、等保测评"],
            ["服务商扰民", "商业推送过度，影响居民体验", "服务商进入小区必须审批，可评价、可投诉、可暂停、可清退"],
            ["试点数据质量差", "房屋、业主、账单和小区关系不准确", "上线前做数据清洗模板，试运营期建立数据纠错机制"],
        ],
        [1700, 3050, 4610],
        font_size=8.6,
    )

    add_heading(doc, "15. 验收标准与交付物", 1)
    add_para(
        doc,
        "验收不应只看页面是否存在，而应看业务链路是否真实闭环、数据是否可追溯、授权是否有效、导出证据是否完整。建议将验收拆成系统能力、业务数据、运营培训和安全配置四类。",
    )
    add_table(
        doc,
        ["类别", "验收内容", "交付物"],
        [
            ["系统环境", "后端、管理端、业主端、数据库、Redis、MinIO 或替代存储可用", "部署说明、环境变量、运行检查报告"],
            ["基础数据", "小区、楼栋、房屋、业主、机构、用户、角色、授权导入完成", "数据导入模板、数据校验表、账号清单"],
            ["业务闭环", "注册准入、缴费、退款、对账、公共收益、审批、投票、工单可跑通", "演示脚本、验收用例、操作截图、导出文件"],
            ["监管能力", "监管大屏、风险预警、信用评分、验收中心、数据交换可用", "监管报告、缺口清单、数据交换包"],
            ["安全审计", "越权拦截、回调验签、防重放、敏感字段审计、导出留痕", "安全体检报告、审计日志、回调诊断记录"],
            ["培训运维", "政府、居委会、物业、银行、业主端使用培训完成", "培训材料、签到记录、FAQ、运维手册"],
        ],
        [1600, 4550, 3210],
        font_size=8.7,
    )

    add_heading(doc, "16. 推广路径：先跑通一个样板，再复制一座城市", 1)
    add_para(
        doc,
        "推广的关键不是一开始铺得多，而是第一批样板要跑得透。先选一个街道、几个居委会、若干类型小区，把账单缴费、公共收益、银行对账、报修投诉、监管预警和验收导出真实跑通；再把这套规则沉淀为模板，复制到更多街道、更多银行、更多物业公司。",
    )
    add_callout(
        doc,
        "推广口径",
        "先拿下样板，再形成标准；先让一个街道看得见，再让一个区县可复制。不要只卖系统，要卖“早一天上线，群众早一天受益；晚一天透明，基层多一天隐患”的治理方案。",
        fill="EDF7F6",
        accent=TEAL,
    )
    add_heading(doc, "16.1 四级推广路径", 2)
    add_numbered(
        doc,
        [
            "样板小区：先选 3-5 个类型不同的小区，把真实账单、缴费、工单、公共收益和监管验收跑通。",
            "街道模板：在一个街道内扩展到 10-30 个小区，形成居委会准入、物业授权和银行协同规则。",
            "区县平台：接入更多街道和银行分支，形成统一监管大屏、信用评价、数据交换和验收口径。",
            "城市运营：引入服务商生态、社区金融、AI 风险预警和多城市渠道合作，把平台从项目变成运营网络。",
        ],
    )
    add_heading(doc, "16.2 立即行动清单", 2)
    add_bullets(
        doc,
        [
            "定名单：明确街道/区县、居委会、物业公司、银行和首批小区，避免项目停在概念层。",
            "定数据：明确小区、楼栋、房屋、业主、账单、银行服务和审批规则，避免上线后没有真实业务。",
            "定接口：明确支付、银行、短信、实名、税务和对象存储先 dev 演示还是直接真实联调。",
            "定指标：明确线上缴费率、工单闭环率、对账差异闭环率、公共收益公示率和监管预警处理率。",
            "定商务：明确试点建设费、物业 SaaS 订阅、银行服务费、实施运维和后续复制价格。",
        ],
    )
    add_heading(doc, "16.3 对外传播词库", 2)
    add_table(
        doc,
        ["对象", "关键词", "一句话"],
        [
            ["政府/街道", "人民、民生、治理", "早一天上线，群众早一天看见透明，基层早一天减少隐患。"],
            ["居委会", "抓手、准入、留痕", "过去靠人盯，现在靠平台留证；治理不能只靠微信群。"],
            ["物业", "增收、省心、少扯皮", "收费更顺，解释更清，服务更有底气，续约更有把握。"],
            ["银行", "入口、资金、场景", "谁先占住小区资金入口，谁先拿到社区金融未来。"],
            ["业主", "透明、安心、能参与", "每一笔钱都看得见，每一次支出都能追，每一次服务都能评。"],
            ["服务商", "合规、评价、长期单", "合规进场才有长期单，不合规只能被挡在小区门外。"],
        ],
        [1700, 2100, 5560],
        font_size=8.7,
    )
    add_image(doc, "宣传片输出/ui-preview-sheet.png", "图 4 已有宣传片抽帧中的功能界面汇总，可用于路演和客户沟通。", width=5.8)

    add_heading(doc, "附录 A. 功能清单", 1)
    add_table(
        doc,
        ["领域", "功能"],
        [
            ["账号权限", "登录 token、角色信息、租户切换、能力菜单、数据授权、详情防越权"],
            ["注册准入", "机构注册、小区备案、物业/银行/服务商绑定、住户实名房屋绑定"],
            ["物业收费", "收费标准、账单生成、在线缴费、退款、渠道账单、欠费台账、导出审计"],
            ["公共收益", "专户管理、支出申请、材料核验、审批模板、业主实名投票、银行放款"],
            ["财务票据", "电子缴费票据、税务开票适配、财务凭证、总账、现金流量表、PDF 导出"],
            ["银行对账", "银行配置、流水导入、对账记录、差异处理、回调验签、放款回执"],
            ["业主服务", "通知公告、报修投诉、问卷投票、多房屋切换、消息回执、服务评价"],
            ["监管治理", "监管大屏、专题分析、风险预警、信用评分、信用因子、数据交换"],
            ["安全运维", "外部接口安全体检、部署就绪检查、敏感字段审计、防重放、验收中心"],
        ],
        [1700, 7660],
        font_size=8.8,
    )

    add_heading(doc, "附录 B. 参考资料与材料来源", 1)
    add_heading(doc, "B.1 本地项目材料", 2)
    add_bullets(
        doc,
        [
            "README.md：工程结构、已实现能力、当前迭代能力和国产信创预留。",
            "多租户注册制落地方案.md：租户、小区关系、数据授权、注册审核和一期排期。",
            "原PDF文本提取.txt：鼎盛阳光物业公共收益管理平台系统方案的政策背景、功能架构和实施路径。",
            "交付物/SPARK_Nexus城市物业治理中枢_商业计划书.md：商业定位、客户角色、收入模式、推广路径和风险应对。",
            "本地截图与宣传片抽帧：管理端、业主端、平台入口和功能展示图。",
        ],
    )
    add_heading(doc, "B.2 联网核验政策资料", 2)
    add_bullets(
        doc,
        [
            "中共中央、国务院：《关于加强基层治理体系和治理能力现代化建设的意见》，2021-07-11，https://www.mee.gov.cn/zcwj/zyygwj/202107/t20210712_846188.shtml",
            "共产党员网：《习近平谈治国理政》第四卷“坚持人民至上”，2022-09-02，https://www.12371.cn/2022/09/02/ARTI1662105164059870.shtml",
            "国务院办公厅：《“十四五”城乡社区服务体系建设规划》，国办发〔2021〕56号，https://www.mee.gov.cn/zcwj/gwywj/202201/t20220124_967972.shtml",
            "九部门：《关于深入推进智慧社区建设的意见》，民发〔2022〕29号，公开资料 https://dsjj.taian.gov.cn/module/download/downfile.jsp?classid=0&filename=50dfd1088d2244f08284b47409c1755d.pdf",
            "住房和城乡建设部办公厅、民政部办公厅：《关于开展完整社区建设试点工作的通知》，公开资料 https://www.hunan.gov.cn/zqt/zcsd/202210/t20221031_29111058.html",
            "《中华人民共和国民法典》第二编物权，业主建筑物区分所有权与共有收益依据，https://www.spp.gov.cn/spp/ssmfdyflvdtpgz/202008/t20200831_478410.shtml",
            "湖北省住房和城乡建设厅：《关于印发〈湖北省住宅小区公共收益管理办法（试行）〉的通知》，2025-07-07，效力状态显示为有效，https://zjt.hubei.gov.cn/zfxxgk/zc/gfxwj/202507/t20250716_5723743.shtml",
            "福建省住房和城乡建设厅等4部门：《福建省住宅小区公共收益管理办法（试行）》，2022-09-30，https://zjt.fujian.gov.cn/xxgk/zfxxgkzl/xxgkml/dfxfgzfgzhgfxwj/gfxwj/202210/t20221021_6021183.htm",
            "河南省住房和城乡建设厅：《河南省住宅小区公共收益管理办法（试行）》，公开资料 https://fczx.pds.gov.cn/contents/16420/466166.html",
            "海南省住房和城乡建设厅：2026-05-27 公开征求《关于规范住宅小区公共收益管理的指导意见（试行）（征求意见稿）》意见，官方列表页 https://zjt.hainan.gov.cn/szjt/0503/nlist2.shtml?ddtab=true；南海网报道 https://m.hinews.cn/page?m=1&n=2826960&s=1044",
            "惠州市住房和城乡建设局：《关于进一步推进开通使用惠州市智慧物业服务平台工作的通知》，https://zjj.huizhou.gov.cn/zwgk/jcgk/content/post_5755557.html",
            "连云港市人民政府：住宅小区公共收益账户共管制度试点，https://www.lyg.gov.cn/zglygzfmhwz/fwzx_zffw/content/81608775-4ba9-458a-b135-8fe22d213374.html",
            "黑龙江省人民政府：七台河智慧社区信息化平台案例，https://www.hlj.gov.cn/hljapp/c116059/202407/c00_31752635.shtml",
            "上海市房屋管理局：“美好家园”典型案例——闵行区新虹街道爱博四村小区，https://fgj.sh.gov.cn/mhjydxal/20211202/ff9b054d5a164b67bf52bf0412597707.html",
            "新华网民生观察：小区公共收益如何不再“隐身”，2026-01-17，https://www.news.cn/local/20260117/e931912b5d8b4033a9934b0432c36b37/c.html",
        ],
    )
    add_callout(
        doc,
        "使用提示",
        "本推介书用于商务沟通、试点申报和项目启动材料。涉及正式采购、合规审查、银行接口和政策条款时，应以最新正式文件、客户制度和接口协议为准。",
        fill="FFF8E8",
        accent=GOLD,
    )

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build_doc()
